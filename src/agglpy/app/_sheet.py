"""The image sheet: one row per image, ``images.xlsx`` or ``images.csv``.

``read_grid`` reads the cells as stored (xlsx values, CSV text);
``validate`` checks them and lists every problem by row and column in
one ``SheetError``. User columns (any column that is not a fixed one)
are kept as stored; ``typed_user_columns`` gives their typed view.
"""

import csv
import io
import math
import re
import unicodedata
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.formula import ArrayFormula

from agglpy.errors import FileFormatError, SheetError

__all__ = [
    "FIXED_COLUMNS",
    "IMAGES_TABLE_COLUMNS",
    "RESERVED_COLUMNS",
    "SHEET_FILES",
    "Entry",
    "Grid",
    "Sheet",
    "find_sheet",
    "read_grid",
    "typed_user_columns",
    "validate",
]

SHEET_FILES = ("images.xlsx", "images.csv")

FIXED_COLUMNS = ("image", "file", "include", "pixel_size_nm", "note")
# Columns agglpy adds to the images table of the results; a user column
# of the same name would clash with them.
IMAGES_TABLE_COLUMNS = (
    "image",
    "pixel_size",
    "pixel_size_source",
    "fov_area",
    "width",
    "height",
    "particle_set",
    "outside_region",
)
RESERVED_COLUMNS = frozenset(FIXED_COLUMNS + IMAGES_TABLE_COLUMNS)

_REQUIRED = ("image", "file")
_PIXEL_SIZE_RANGE = (0.001, 10_000.0)
_TRUE = {"true", "yes", "1"}
_FALSE = {"false", "no", "0"}
# Not allowed in a folder name on Windows (the strictest system).
_BAD_CHARACTERS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_DEVICE_NAMES = re.compile(r"(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])", re.I)
_NUMERIC_TEXT = re.compile(r"[+-]?(0|[1-9][0-9]*)(\.[0-9]+)?")


def find_sheet(folder: Path) -> Path:
    """The image sheet in ``folder``: ``images.xlsx`` or ``images.csv``.

    Raises SheetError when both exist, FileNotFoundError when neither.
    """
    found = [folder / name for name in SHEET_FILES if (folder / name).exists()]
    if len(found) > 1:
        raise SheetError(
            f"{folder} has both {SHEET_FILES[0]} and {SHEET_FILES[1]}; "
            "keep one"
        )
    if not found:
        raise FileNotFoundError(
            f"{folder} has no {SHEET_FILES[0]} or {SHEET_FILES[1]}; "
            "create a collection with Collection.create"
        )
    return found[0]


@dataclass
class Grid:
    """The cells of a sheet as stored.

    Attributes:
        path: The sheet file.
        header: The column names (row 1).
        rows: ``(row number, cells)`` for every row that is not fully
            empty; row numbers as Excel shows them (header = 1); cells
            None when empty, xlsx values or CSV text.
    """

    path: Path
    header: list[str]
    rows: list[tuple[int, list[object]]]


@dataclass(frozen=True)
class Entry:
    """One image row of a valid sheet.

    Attributes:
        row: Its row number in the sheet.
        image: The image name.
        file: The image file as written in the sheet.
        include: Whether the image is analysed.
        pixel_size_nm: The pixel size in nm, or None (taken from the
            image file then).
        note: Free text.
        user: The user columns' cells as stored.
    """

    row: int
    image: str
    file: str
    include: bool
    pixel_size_nm: float | None
    note: str
    user: Mapping[str, object]


@dataclass(frozen=True)
class Sheet:
    """A valid image sheet.

    Attributes:
        path: The sheet file.
        columns: All column names in file order.
        entries: The image rows in file order.
    """

    path: Path
    columns: tuple[str, ...]
    entries: tuple[Entry, ...]

    @property
    def user_columns(self) -> tuple[str, ...]:
        """The columns that are not fixed ones, in file order."""
        return tuple(c for c in self.columns if c not in FIXED_COLUMNS)

    @property
    def names(self) -> tuple[str, ...]:
        """The image names in file order."""
        return tuple(e.image for e in self.entries)


# ---------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------


def read_grid(path: Path) -> Grid:
    """Read the cells of an image sheet.

    Raises FileFormatError when the file cannot be read and SheetError
    for formulas, more than one worksheet or a bad header.
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        table = _read_csv(path)
    else:
        table = _read_xlsx(path)
    return _grid(path, table)


def _read_csv(path: Path) -> list[tuple[int, list[object]]]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        raise FileFormatError(
            f"{path} is not UTF-8 text; save it as CSV UTF-8 or use "
            "images.xlsx"
        ) from None
    reader = csv.reader(io.StringIO(text, newline=""))
    return [
        (number, [cell if cell != "" else None for cell in cells])
        for number, cells in enumerate(reader, start=1)
    ]


def _read_xlsx(path: Path) -> list[tuple[int, list[object]]]:
    try:
        book = openpyxl.load_workbook(path)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as e:
        raise FileFormatError(
            f"{path} is not a readable xlsx file ({e})"
        ) from None
    if len(book.worksheets) > 1:
        names = ", ".join(s.title for s in book.worksheets)
        raise SheetError(
            f"{path.name} has {len(book.worksheets)} worksheets ({names}); "
            "keep one"
        )
    sheet = book.worksheets[0]
    formulas: list[tuple[int, int]] = []
    table = []
    # Rows and columns counted from 1, as Excel shows them.
    for number, cells in enumerate(sheet.iter_rows(min_row=1), start=1):
        values: list[object] = []
        for column, cell in enumerate(cells, start=1):
            if cell.data_type == "f" or isinstance(cell.value, ArrayFormula):
                formulas.append((column, number))
            value = cell.value
            values.append(None if value == "" else value)
        table.append((number, values))
    if formulas:
        raise SheetError(_formula_message(path, formulas))
    return table


def _formula_message(path: Path, cells: list[tuple[int, int]]) -> str:
    """Formula cells as contiguous runs per column, e.g. ``D2:D3, F4``."""
    runs: list[str] = []
    for column in sorted({c for c, _ in cells}):
        rows = sorted(r for c, r in cells if c == column)
        letter = get_column_letter(column)
        start = prev = rows[0]
        for r in [*rows[1:], None]:
            if r is not None and r == prev + 1:
                prev = r
                continue
            runs.append(
                f"{letter}{start}"
                if start == prev
                else f"{letter}{start}:{letter}{prev}"
            )
            if r is not None:
                start = prev = r
    if len(cells) == 1:
        what = "contains a formula; replace it with a value"
    else:
        what = "contain formulas; replace them with values"
    return f"{path.name}: {', '.join(runs)} {what} (Paste Special → Values)"


def _empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and value == "")


def _grid(path: Path, table: list[tuple[int, list[object]]]) -> Grid:
    rows = [(n, cells) for n, cells in table if not all(map(_empty, cells))]
    header_cells: list[object] = []
    if rows and rows[0][0] == 1:
        header_cells = rows.pop(0)[1]
    width = max([len(header_cells), *(len(c) for _, c in rows)])
    has_data = [
        any(i < len(c) and not _empty(c[i]) for _, c in rows)
        for i in range(width)
    ]
    header_cells = header_cells + [None] * (width - len(header_cells))
    # Empty header cells at the right with nothing under them are not
    # columns (Excel keeps formatted empty cells).
    while header_cells and _empty(header_cells[-1]) and not has_data[-1]:
        header_cells.pop()
        has_data.pop()
    width = len(header_cells)

    problems: list[str] = []
    header: list[str] = []
    for i, name in enumerate(header_cells):
        letter = get_column_letter(i + 1)
        if _empty(name):
            if has_data[i]:
                problems.append(
                    f"column {letter}: no name in row 1, but the column "
                    "has values"
                )
            header.append("")
        elif not isinstance(name, str):
            problems.append(f"column {letter}: the name {name!r} is not text")
            header.append(str(name))
        else:
            header.append(name)
    for name in dict.fromkeys(header):
        where = [
            get_column_letter(i + 1) for i, h in enumerate(header) if h == name
        ]
        if name and len(where) > 1:
            problems.append(
                f"column {name!r} is repeated (columns {', '.join(where)})"
            )
    _raise(path, problems)
    cells = [(n, (c + [None] * width)[:width]) for n, c in rows]
    return Grid(path, header, cells)


# ---------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------


def validate(grid: Grid) -> Sheet:
    """Check the sheet's cells; raise SheetError listing every problem.

    ``pixel_size_nm`` is in nm (0.001 to 10 000).
    """
    problems: list[tuple[int, str]] = []
    header = grid.header
    for name in _REQUIRED:
        if name not in header:
            similar = [h for h in header if h.casefold() == name]
            hint = (
                f" (column names are case-sensitive; found {similar[0]!r})"
                if similar
                else ""
            )
            problems.append((1, f"column {name!r} is missing{hint}"))
    for name in header:
        if name in RESERVED_COLUMNS and name not in FIXED_COLUMNS:
            problems.append(
                (
                    1,
                    f"column {name!r}: agglpy adds a column of this name to "
                    "the images table; rename it",
                )
            )
    if any(name not in header for name in _REQUIRED):
        _raise(grid.path, [text for _, text in problems])

    index = {name: i for i, name in enumerate(header)}
    entries = []
    for number, cells in grid.rows:

        def cell(name: str, cells: list[object] = cells) -> object:
            return cells[index[name]] if name in index else None

        entry, found = _entry(number, cell, header, index)
        problems.extend((number, f"row {number}, {text}") for text in found)
        if entry is not None:
            entries.append(entry)
    problems.extend(_duplicates(entries, grid.rows, index))
    problems.sort(key=lambda p: p[0])
    _raise(grid.path, [text for _, text in problems])
    return Sheet(grid.path, tuple(header), tuple(entries))


def _entry(
    number: int,
    cell: Any,
    header: list[str],
    index: Mapping[str, int],
) -> tuple[Entry | None, list[str]]:
    found: list[str] = []
    image = "" if _empty(cell("image")) else _text(cell("image"))
    bad_name = _name_problem(image)
    if bad_name:
        found.append(f"image: {bad_name}")
    file = cell("file")
    if _empty(file):
        found.append("file: required")
    include = _include(cell("include"))
    if include is None:
        found.append(
            f"include: {cell('include')!r} is not true or false "
            "(accepted: TRUE/FALSE cells, true/false, yes/no, 1/0; empty = "
            "included). A CSV saved by Excel in Polish writes PRAWDA/FAŁSZ; "
            "images.xlsx keeps TRUE/FALSE cells"
        )
    pixel_size, problem = _pixel_size(cell("pixel_size_nm"))
    if problem:
        found.append(f"pixel_size_nm: {problem}")
    if found:
        return None, found
    note = cell("note")
    user = {name: cell(name) for name in header if name not in FIXED_COLUMNS}
    entry = Entry(
        row=number,
        image=image,
        file=_text(file),
        include=bool(include),
        pixel_size_nm=pixel_size,
        note="" if _empty(note) else _text(note),
        user=user,
    )
    return entry, []


def _name_problem(name: str) -> str | None:
    if not name:
        return "required"
    if _BAD_CHARACTERS.search(name):
        return (
            f"{name!r} can't be a folder name "
            '(no < > : " / \\ | ? * or control characters)'
        )
    if name[-1] in ". ":
        return f"{name!r} can't end with a dot or a space (Windows)"
    if _DEVICE_NAMES.fullmatch(name.split(".")[0]):
        return f"{name!r} is a reserved device name on Windows"
    return None


def _duplicates(
    entries: list[Entry],
    rows: list[tuple[int, list[object]]],
    index: Mapping[str, int],
) -> list[tuple[int, str]]:
    """Names repeated ignoring case, over every row with a name."""
    by_key: dict[str, list[tuple[int, str]]] = {}
    for number, cells in rows:
        value = cells[index["image"]]
        if _empty(value):
            continue
        name = _text(value)
        key = unicodedata.normalize("NFC", name).casefold()
        by_key.setdefault(key, []).append((number, name))
    problems = []
    for found in by_key.values():
        if len(found) < 2:
            continue
        numbers = ", ".join(str(n) for n, _ in found)
        names = list(dict.fromkeys(name for _, name in found))
        if len(names) == 1:
            text = f"{names[0]!r} is repeated"
        else:
            shown = [repr(n) for n in names]
            text = (
                f"{', '.join(shown[:-1])} and {shown[-1]} differ only in case"
            )
        problems.append((found[0][0], f"rows {numbers}, image: {text}"))
    return problems


def _include(value: object) -> bool | None:
    """True / False, or None when the value is not accepted."""
    if _empty(value):
        return True
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        word = value.strip().casefold()
        if word in _TRUE:
            return True
        if word in _FALSE:
            return False
    return None


def _pixel_size(value: object) -> tuple[float | None, str | None]:
    if _empty(value):
        return None, None
    number: float | None = None
    if isinstance(value, bool):
        number = None
    elif isinstance(value, int | float):
        number = float(value)
    elif isinstance(value, str) and _NUMERIC_TEXT.fullmatch(value.strip()):
        number = float(value)
    if number is None or not math.isfinite(number):
        hint = (
            " (use . as the decimal point)"
            if isinstance(value, str) and "," in value
            else ""
        )
        return None, f"{value!r} is not a number{hint}"
    low, high = _PIXEL_SIZE_RANGE
    if not low <= number <= high:
        text = f"{number:g} is outside {low:g} to {high:g} nm; the column is in nm"
        if 0 < number < low:
            text += f" ({number:g} m = {number * 1e9:g} nm)"
        return None, text
    return number, None


def _text(value: object) -> str:
    """A cell as text: whole floats without ``.0``, bools TRUE / FALSE."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)


def _raise(path: Path, problems: Sequence[str]) -> None:
    if not problems:
        return
    noun = "problem" if len(problems) == 1 else "problems"
    lines = "\n".join(f"  {p}" for p in problems)
    raise SheetError(f"{path.name} has {len(problems)} {noun}:\n{lines}")


# ---------------------------------------------------------------------
# Typed view
# ---------------------------------------------------------------------


def typed_user_columns(sheet: Sheet) -> pd.DataFrame:
    """Column ``image`` and the user columns, typed.

    Per column, over its non-empty cells: all bools → boolean; all
    numbers or numeric text without a leading zero → int64 when all are
    whole and none is empty, else float64; all dates → datetime64;
    otherwise text (so IDs such as ``007`` stay text).
    """
    data: dict[str, Any] = {"image": [e.image for e in sheet.entries]}
    for name in sheet.user_columns:
        data[name] = _typed([e.user[name] for e in sheet.entries])
    return pd.DataFrame(data)


def _typed(values: list[object]) -> pd.Series:
    present = [v for v in values if not _empty(v)]
    if present and all(isinstance(v, bool) for v in present):
        return pd.Series(
            [pd.NA if _empty(v) else v for v in values], dtype="boolean"
        )
    numbers = [_number(v) for v in present]
    if present and all(n is not None for n in numbers):
        whole = all(float(n).is_integer() for n in numbers if n is not None)
        if whole and len(present) == len(values):
            return pd.Series(
                [int(_number(v) or 0) for v in values], dtype="int64"
            )
        return pd.Series(
            [math.nan if _empty(v) else _number(v) for v in values],
            dtype="float64",
        )
    if present and all(isinstance(v, datetime | date) for v in present):
        cells = pd.Series([None if _empty(v) else v for v in values])
        return pd.to_datetime(cells)
    return pd.Series(
        [None if _empty(v) else _text(v) for v in values], dtype="object"
    )


def _number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, str) and _NUMERIC_TEXT.fullmatch(value):
        return float(value)
    return None
