"""Write the image sheet: in place, or a new one.

An existing sheet is edited, never rebuilt: only changed cells are set,
rows removed or appended, so a person's formatting (widths, colours,
filters, frozen rows) stays. Rows are matched by ``image``.
"""

import csv
import io
import math
import numbers
from collections.abc import Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from agglpy.app._files import replace_file
from agglpy.app._sheet import read_grid

__all__ = ["new_sheet", "write_sheet"]

_BOM = b"\xef\xbb\xbf"
_TRUE = {"true", "yes", "1"}
_FALSE = {"false", "no", "0"}


def write_sheet(
    path: Path, header: Sequence[str], rows: Sequence[Sequence[object]]
) -> None:
    """Make the sheet at ``path`` hold ``header`` and ``rows``.

    Rows are matched by ``image`` (ignoring case): changed cells are
    set, missing rows removed, new rows appended in the order given.
    New columns are added at the right, missing ones deleted. The
    file's row order is kept.
    """
    header = list(header)
    cells = [[_plain(v) for v in row] for row in rows]
    if path.suffix.lower() == ".csv":
        _write_csv(path, header, cells)
    else:
        _write_xlsx(path, header, cells)


def new_sheet(
    path: Path, header: Sequence[str], rows: Sequence[Sequence[object]]
) -> None:
    """Write a new sheet: xlsx with a bold, frozen header, or CSV."""
    cells = [[_plain(v) for v in row] for row in rows]
    if path.suffix.lower() == ".csv":
        text = _csv_text(
            [list(header), *([_text(v) for v in r] for r in cells)]
        )
        replace_file(path, lambda tmp: tmp.write_bytes(text.encode()))
        return
    book = openpyxl.Workbook()
    sheet = book.worksheets[0]
    sheet.title = "images"
    sheet.append(list(header))
    for row in cells:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    replace_file(path, book.save)


class _Plan:
    """What changes between the sheet on disk and the rows wanted."""

    def __init__(
        self,
        old_header: list[str],
        old_rows: list[tuple[int, list[object]]],
        header: list[str],
        rows: list[list[Any]],
    ) -> None:
        key = header.index("image")
        old_key = old_header.index("image")
        wanted = {_key(row[key]): row for row in rows}
        on_disk = {_key(cells[old_key]) for _, cells in old_rows}
        # Columns: kept ones stay where they are; unnamed ones (spacers)
        # are left alone.
        self.removed_columns = [
            i
            for i, name in enumerate(old_header)
            if name and name not in header
        ]
        self.columns = [n for n in old_header if not n or n in header]
        self.columns += [n for n in header if n not in old_header]
        self.removed_rows = [
            number
            for number, cells in old_rows
            if _key(cells[old_key]) not in wanted
        ]
        self.kept = [
            (number, wanted[_key(cells[old_key])], cells)
            for number, cells in old_rows
            if _key(cells[old_key]) in wanted
        ]
        self.new_rows = [row for row in rows if _key(row[key]) not in on_disk]
        self.header = header
        self.old_header = old_header

    def value(self, row: list[Any], column: str) -> Any:
        return row[self.header.index(column)]


def _write_xlsx(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    grid = read_grid(path)
    plan = _Plan(grid.header, grid.rows, header, rows)
    book = openpyxl.load_workbook(path)
    sheet = book.worksheets[0]

    for i in sorted(plan.removed_columns, reverse=True):
        sheet.delete_cols(i + 1)
        _remove_table_column(sheet, i + 1)
    for number in sorted(plan.removed_rows, reverse=True):
        sheet.delete_rows(number)
    position = {name: i + 1 for i, name in enumerate(plan.columns) if name}
    for name in plan.columns:
        if name and name not in grid.header:
            sheet.cell(row=1, column=position[name], value=name)

    def shifted(number: int) -> int:
        return number - sum(1 for r in plan.removed_rows if r < number)

    last = 1
    for number, row, old in plan.kept:
        target = shifted(number)
        last = max(last, target)
        for name in header:
            old_value = (
                old[grid.header.index(name)] if name in grid.header else None
            )
            new_value = plan.value(row, name)
            if not _same(old_value, new_value):
                # cell(value=None) would leave the old value.
                cell = sheet.cell(row=target, column=position[name])
                cell.value = new_value
    for row in plan.new_rows:
        # Explicit row numbers: formatted empty rows below the data
        # would make ``append`` start further down.
        last += 1
        for name in header:
            value = plan.value(row, name)
            if value is not None:
                sheet.cell(row=last, column=position[name], value=value)
    _resize_ranges(sheet, last)
    replace_file(path, book.save)


def _remove_table_column(sheet: Worksheet, column: int) -> None:
    """Shrink the Excel tables and the filter over a deleted column
    (openpyxl does not; Excel would report the file as damaged)."""
    for table in sheet.tables.values():
        min_col, _, max_col, _ = range_boundaries(table.ref)
        if min_col is not None and min_col <= column <= (max_col or 0):
            del table.tableColumns[column - min_col]
        table.ref = _without_column(table.ref, column)
        if table.autoFilter is not None and table.autoFilter.ref:
            table.autoFilter.ref = _without_column(
                table.autoFilter.ref, column
            )
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = _without_column(sheet.auto_filter.ref, column)


def _without_column(ref: str, column: int) -> str:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    if min_col is None or min_row is None or max_col is None:
        return ref
    if column < min_col:
        min_col -= 1
    if column <= max_col:
        max_col -= 1
    return (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )


def _resize_ranges(sheet: Worksheet, last: int) -> None:
    """Let Excel tables and the filter that start at the header end at
    the last data row (openpyxl does not move them)."""
    last = max(last, 2)  # an Excel table needs one data row
    for table in sheet.tables.values():
        table.ref = _resized(table.ref, last)
        if table.autoFilter is not None and table.autoFilter.ref:
            table.autoFilter.ref = _resized(table.autoFilter.ref, last)
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = _resized(sheet.auto_filter.ref, last)


def _resized(ref: str, last: int) -> str:
    min_col, min_row, max_col, _ = range_boundaries(ref)
    if min_row != 1 or min_col is None or max_col is None:
        return ref
    return f"{get_column_letter(min_col)}1:{get_column_letter(max_col)}{last}"


def _write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig")
    newline = "\r\n" if "\r\n" in text else "\n"
    records = list(csv.reader(io.StringIO(text, newline="")))
    grid = read_grid(path)
    plan = _Plan(grid.header, grid.rows, header, rows)
    by_number = {number: (row, old) for number, row, old in plan.kept}
    removed = set(plan.removed_rows)
    width = len(grid.header)

    out: list[list[str]] = [list(plan.columns)]
    for number, record in enumerate(records[1:], start=2):
        if number in removed:
            continue
        cells = (record + [""] * width)[: max(width, len(record))]
        if number not in by_number:
            # A fully empty line stays as it was.
            out.append(record)
            continue
        row, _ = by_number[number]
        line = []
        for name in plan.columns:
            if name in grid.header:
                old_text = cells[grid.header.index(name)]
            else:
                old_text = ""
            if not name:
                line.append(old_text)
                continue
            new_value = plan.value(row, name)
            old = None if old_text == "" else old_text
            line.append(
                old_text if _same(old, new_value) else _text(new_value)
            )
        out.append(line)
    for row in plan.new_rows:
        out.append(
            ["" if not n else _text(plan.value(row, n)) for n in plan.columns]
        )
    encoded = (_BOM if raw.startswith(_BOM) else b"") + _csv_text(
        out, newline
    ).encode()
    replace_file(path, lambda tmp: tmp.write_bytes(encoded))


def _csv_text(records: list[list[str]], newline: str = "\n") -> str:
    stream = io.StringIO()
    writer = csv.writer(stream, lineterminator=newline)
    writer.writerows(records)
    return stream.getvalue()


def _key(value: object) -> str:
    return _text(value).casefold()


def _plain(value: object) -> Any:
    """A cell value from pandas / NumPy as a plain Python value."""
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _same(old: object, new: object) -> bool:
    """Whether a stored cell already holds the wanted value."""
    if old is None or new is None:
        return old is None and new is None
    if isinstance(new, bool) or isinstance(old, bool):
        if isinstance(old, str):
            word = old.strip().casefold()
            return (word in _TRUE and new is True) or (
                word in _FALSE and new is False
            )
        return isinstance(old, bool) and isinstance(new, bool) and old == new
    if isinstance(new, numbers.Real):
        return _as_number(old) == float(new)
    return _text(old) == _text(new) if isinstance(new, str) else old == new


def _as_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, numbers.Real):
        return float(value)
    if isinstance(value, str):
        try:
            number = float(value)
        except ValueError:
            return None
        return number if math.isfinite(number) else None
    return None


def _text(value: object) -> str:
    """A value as CSV text: whole floats without ``.0``, bools true/false."""
    value = _plain(value)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)
