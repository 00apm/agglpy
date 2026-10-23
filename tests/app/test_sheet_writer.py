"""Writing the image sheet back in place, and safe file replacement."""

import os
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table

from agglpy.app import _files
from agglpy.app._files import replace_file
from agglpy.app._sheet_writer import new_sheet, write_sheet
from agglpy.errors import OutputLockedError

from support.sheets import write_csv, write_xlsx

H = ["image", "file", "include", "pixel_size_nm", "note", "T"]
ROWS = [
    ["a", "raw/a.tif", None, None, None, 20],
    ["b", "raw/b.tif", False, 3.9, "blurry", 100],
    ["c", "raw/c.tif", None, None, None, 7],
]


def cells(path: Path) -> list[list[object]]:
    sheet = openpyxl.load_workbook(path).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


@pytest.fixture
def book(tmp_path: Path) -> Path:
    return write_xlsx(tmp_path / "images.xlsx", [H, *ROWS])


def test_changed_cell_only(book: Path) -> None:
    wb = openpyxl.load_workbook(book)
    ws = wb.active
    ws["B3"].fill = PatternFill("solid", start_color="FFFF00")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["B"].width = 40
    ws.freeze_panes = "A2"
    wb.save(book)
    before = cells(book)

    rows = [list(r) for r in ROWS]
    rows[1][4] = "sharp"
    write_sheet(book, H, rows)

    after = cells(book)
    changed = [
        (r, c)
        for r, (old, new) in enumerate(zip(before, after, strict=True))
        for c, (x, y) in enumerate(zip(old, new, strict=True))
        if x != y
    ]
    assert changed == [(2, 4)]
    assert after[2][4] == "sharp"
    ws = openpyxl.load_workbook(book).active
    assert ws["B3"].fill.start_color.rgb.endswith("FFFF00")
    assert ws["A1"].font.bold
    assert ws.column_dimensions["B"].width == 40
    assert ws.freeze_panes == "A2"


def test_equal_number_is_unchanged(book: Path) -> None:
    rows = [list(r) for r in ROWS]
    rows[0][5] = 20.0
    write_sheet(book, H, rows)
    value = openpyxl.load_workbook(book).active["F2"].value
    assert value == 20
    assert type(value) is int


def test_rows_removed_and_appended(book: Path) -> None:
    d = ["d", "raw/d.tif", None, None, None, 1]
    write_sheet(book, H, [ROWS[1], d])
    assert cells(book) == [H, ROWS[1], d]


def test_append_after_formatted_empty_rows(book: Path) -> None:
    wb = openpyxl.load_workbook(book)
    ws = wb.active
    for r in range(5, 21):
        ws.cell(row=r, column=1).fill = PatternFill(
            "solid", start_color="DDDDDD"
        )
    wb.save(book)
    d = ["d", "raw/d.tif", None, None, None, 1]
    write_sheet(book, H, [*ROWS, d])
    assert openpyxl.load_workbook(book).active["A5"].value == "d"


def test_append_extends_excel_table_and_filter(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "images.xlsx", [H, *ROWS[:2]])
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.add_table(Table(displayName="T1", ref="A1:F3"))
    ws.auto_filter.ref = "A1:F3"
    wb.save(path)
    write_sheet(path, H, ROWS)
    ws = openpyxl.load_workbook(path).active
    assert ws.tables["T1"].ref == "A1:F4"
    assert ws.auto_filter.ref == "A1:F4"


def test_removed_row_shrinks_excel_table(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "images.xlsx", [H, *ROWS])
    wb = openpyxl.load_workbook(path)
    wb.active.add_table(Table(displayName="T1", ref="A1:F4"))
    wb.save(path)
    write_sheet(path, H, ROWS[:2])
    assert openpyxl.load_workbook(path).active.tables["T1"].ref == "A1:F3"


def test_new_column_appended_and_removed_column_deleted(book: Path) -> None:
    header = ["image", "file", "include", "pixel_size_nm", "T", "series"]
    rows = [[*r[:4], r[5], "D7"] for r in ROWS]
    write_sheet(book, header, rows)
    assert cells(book) == [header, *rows]


def test_file_row_order_kept(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "images.xlsx", [H, ROWS[0], ROWS[1]])
    write_sheet(path, H, [ROWS[1], ROWS[0], ROWS[2]])
    assert cells(path) == [H, *ROWS]


def test_rows_matched_by_image_ignoring_case(book: Path) -> None:
    rows = [list(r) for r in ROWS]
    rows[0][0] = "A"
    write_sheet(book, H, rows)
    assert cells(book)[1][0] == "A"
    assert len(cells(book)) == 4


def test_pandas_values_normalised(book: Path) -> None:
    rows = [list(r) for r in ROWS]
    rows[0][5] = np.int64(3)
    rows[1][3] = np.nan
    rows[2][5] = pd.NA
    write_sheet(book, H, rows)
    after = cells(book)
    assert (after[1][5], after[2][3], after[3][5]) == (3, None, None)


# ---------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------


def csv_rows() -> list[list[object]]:
    return [
        H,
        ["a", "raw/a.tif", "", "", "", "007"],
        ["b", "raw/b.tif", "no", "1.50", "", "20"],
    ]


def test_csv_unchanged_text_kept(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "images.csv", csv_rows())
    write_sheet(
        path,
        H,
        [
            ["a", "raw/a.tif", None, None, "new note", "007"],
            ["b", "raw/b.tif", False, 1.5, None, 20],
        ],
    )
    assert path.read_text(encoding="utf-8") == (
        "image,file,include,pixel_size_nm,note,T\n"
        "a,raw/a.tif,,,new note,007\n"
        "b,raw/b.tif,no,1.50,,20\n"
    )


def test_csv_bom_and_line_ends_kept(tmp_path: Path) -> None:
    path = write_csv(
        tmp_path / "images.csv", csv_rows(), bom=True, newline="\r\n"
    )
    write_sheet(path, H, [["a", "raw/a.tif", None, None, None, "007"]])
    assert path.read_bytes() == (
        b"\xef\xbb\xbfimage,file,include,pixel_size_nm,note,T\r\n"
        b"a,raw/a.tif,,,,007\r\n"
    )


def test_csv_new_values_as_text(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "images.csv", [H])
    write_sheet(
        path,
        H,
        [
            ["a", "raw/a.tif", True, 20.0, None, 0.25],
            ["b", "raw/b.tif", False, np.nan, pd.NA, np.int64(7)],
        ],
    )
    assert path.read_text(encoding="utf-8").splitlines()[1:] == [
        "a,raw/a.tif,true,20,,0.25",
        "b,raw/b.tif,false,,,7",
    ]


# ---------------------------------------------------------------------
# New sheets
# ---------------------------------------------------------------------


def test_new_xlsx_sheet(tmp_path: Path) -> None:
    path = tmp_path / "images.xlsx"
    new_sheet(path, H, ROWS)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["images"]
    ws = wb.active
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert cells(path) == [H, *ROWS]


def test_new_csv_sheet_has_no_bom(tmp_path: Path) -> None:
    path = tmp_path / "images.csv"
    new_sheet(path, H, ROWS[:1])
    assert path.read_bytes() == (
        b"image,file,include,pixel_size_nm,note,T\na,raw/a.tif,,,,20\n"
    )


# ---------------------------------------------------------------------
# replace_file
# ---------------------------------------------------------------------


def temp_files(folder: Path) -> list[str]:
    return sorted(p.name for p in folder.iterdir() if "agglpy-tmp" in p.name)


def test_locked_file_raises_and_keeps_original(
    book: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original = book.read_bytes()

    def locked(src: object, dst: object) -> None:
        raise PermissionError(13, "in use")

    monkeypatch.setattr(_files.os, "replace", locked)
    with pytest.raises(OutputLockedError) as e:
        write_sheet(book, H, ROWS[:1])
    assert "close images.xlsx and run again" in str(e.value)
    assert book.read_bytes() == original
    assert temp_files(book.parent) == []


def test_failed_write_leaves_no_temp_file(tmp_path: Path) -> None:
    path = tmp_path / "agglpy.yml"
    path.write_text("a: 1\n", encoding="utf-8")

    def write(tmp: Path) -> None:
        tmp.write_text("half", encoding="utf-8")
        raise ValueError("boom")

    with pytest.raises(ValueError, match="boom"):
        replace_file(path, write)
    assert path.read_text(encoding="utf-8") == "a: 1\n"
    assert temp_files(tmp_path) == []


def test_replace_file_writes_new_file(tmp_path: Path) -> None:
    path = tmp_path / "agglpy.yml"
    replace_file(path, lambda tmp: tmp.write_text("x", encoding="utf-8"))
    assert path.read_text(encoding="utf-8") == "x"
    assert os.listdir(tmp_path) == ["agglpy.yml"]


def test_removed_column_shrinks_excel_table(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "images.xlsx", [H, *ROWS])
    wb = openpyxl.load_workbook(path)
    wb.active.add_table(Table(displayName="T1", ref="A1:F4"))
    wb.active.auto_filter.ref = "A1:F4"
    wb.save(path)
    header = [h for h in H if h != "note"]
    write_sheet(path, header, [[*r[:4], r[5]] for r in ROWS])
    sheet = openpyxl.load_workbook(path).active
    table = sheet.tables["T1"]
    assert table.ref == "A1:E4"
    assert [c.name for c in table.tableColumns] == header
    assert sheet.auto_filter.ref == "A1:E4"
