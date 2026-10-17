"""agglpy.export.write_results: xlsx, CSV and the record in a folder."""

import dataclasses
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest
import yaml

from agglpy import export
from agglpy.errors import (
    OutputExistsError,
    OutputLockedError,
    ParamsError,
    TableError,
)
from agglpy.params import HCTParams, HCTRange

PARTICLES = pd.DataFrame(
    {
        "id": [1, 2, 3],
        "x": [0.1, 1 / 3, 2e-7],
        "D": [np.nan, 2.0, 1e6],
        "type": ["a", "b", "ä"],
        "touches_edge": [True, False, False],
    }
)
SUMMARY = pd.DataFrame({"N": [3, 4]}, index=pd.Index(["A", "B"], name="group"))


def write(folder: Path, **options) -> Path:
    tables = options.pop(
        "tables", {"particles": PARTICLES, "summary": SUMMARY}
    )
    return export.write_results(folder, tables, **options)


def sheets(folder: Path) -> list[str]:
    return openpyxl.load_workbook(folder / "results.xlsx").sheetnames


def sheet_rows(folder: Path, name: str) -> list[tuple]:
    sheet = openpyxl.load_workbook(folder / "results.xlsx")[name]
    return list(sheet.iter_rows(values_only=True))


def record(folder: Path) -> dict:
    with (folder / "record.yml").open(encoding="utf-8") as fh:
        return yaml.safe_load(fh)


# ---------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------


def test_files_are_written_into_a_new_folder(tmp_path):
    out = write(tmp_path / "a" / "b")
    assert out == tmp_path / "a" / "b"
    assert sorted(p.name for p in out.iterdir()) == [
        "csv",
        "record.yml",
        "results.xlsx",
    ]
    assert sorted(p.name for p in (out / "csv").iterdir()) == [
        "particles.csv",
        "summary.csv",
    ]


def test_sheets_in_the_given_order_then_info(tmp_path):
    out = write(tmp_path, tables={"summary": SUMMARY, "particles": PARTICLES})
    assert sheets(out) == ["summary", "particles", "info"]


def test_csv_reads_back_equal(tmp_path):
    out = write(tmp_path)
    back = pd.read_csv(out / "csv" / "particles.csv", encoding="utf-8")
    pd.testing.assert_frame_equal(back, PARTICLES)
    text = (out / "csv" / "particles.csv").read_bytes()
    assert text.startswith(b"id,x,D,type,touches_edge\n")  # no BOM, LF


def test_xlsx_sheet_holds_the_table(tmp_path):
    out = write(tmp_path)
    rows = sheet_rows(out, "particles")
    assert rows[0] == ("id", "x", "D", "type", "touches_edge")
    assert rows[2] == (2, 1 / 3, 2.0, "b", False)
    assert rows[1][2] is None  # NaN is an empty cell


def test_an_index_with_data_becomes_columns(tmp_path):
    out = write(tmp_path)
    assert (out / "csv" / "summary.csv").read_text() == "group,N\nA,3\nB,4\n"
    assert sheet_rows(out, "summary")[0] == ("group", "N")


def test_a_filtered_table_keeps_its_row_labels(tmp_path):
    out = write(tmp_path, tables={"big": PARTICLES[PARTICLES["id"] > 1]})
    back = pd.read_csv(out / "csv" / "big.csv")
    assert list(back["index"]) == [1, 2]


def test_a_table_over_the_row_limit_is_csv_only(tmp_path, monkeypatch):
    monkeypatch.setattr(export, "_XLSX_ROWS", 2)
    out = write(tmp_path)
    assert sheets(out) == ["summary", "info"]
    assert (out / "csv" / "particles.csv").exists()
    tables = record(out)["agglpy"]["tables"]
    assert tables[0] == {"name": "particles", "rows": 3, "written": "csv only"}
    assert ("agglpy.tables.0.written", "csv only") in sheet_rows(out, "info")


@pytest.mark.parametrize(
    "name", ["", "a b", "ä", "x" * 32, "info", "Units", "a/b", 3]
)
def test_table_names_must_be_safe(tmp_path, name):
    with pytest.raises(TableError, match="name"):
        write(tmp_path, tables={name: PARTICLES})
    assert not (tmp_path / "results.xlsx").exists()


def test_table_names_differ_by_more_than_case(tmp_path):
    with pytest.raises(TableError, match="twice"):
        write(tmp_path, tables={"p": PARTICLES, "P": PARTICLES})


def test_a_table_must_be_a_dataframe(tmp_path):
    with pytest.raises(TableError, match="DataFrame"):
        write(tmp_path, tables={"p": PARTICLES["x"]})


# ---------------------------------------------------------------------
# Units
# ---------------------------------------------------------------------


def test_units_sheet_and_csv_list_every_column_with_a_unit(tmp_path):
    out = write(tmp_path, units={"x": "px", "D": "px", "N": "1", "other": "m"})
    expected = [
        ("table", "column", "unit"),
        ("particles", "x", "px"),
        ("particles", "D", "px"),
        ("summary", "N", "1"),
    ]
    assert sheets(out) == ["particles", "summary", "units", "info"]
    assert sheet_rows(out, "units") == expected
    lines = (out / "csv" / "units.csv").read_text().splitlines()
    assert lines == [",".join(row) for row in expected]


def test_without_units_no_units_sheet(tmp_path):
    out = write(tmp_path)
    assert "units" not in sheets(out)
    assert not (out / "csv" / "units.csv").exists()


def test_a_unit_must_be_text(tmp_path):
    with pytest.raises(TypeError, match="'D'"):
        write(tmp_path, units={"D": 1e-6})


# ---------------------------------------------------------------------
# Record
# ---------------------------------------------------------------------


def test_record_holds_the_callers_entries_and_agglpys(tmp_path):
    params = HCTParams(ranges=(HCTRange(3, 60),))
    out = write(
        tmp_path,
        record={"image": "D7-017", "params": params, "n": np.int64(3)},
    )
    saved = record(out)
    assert list(saved) == ["agglpy", "image", "params", "n"]
    assert saved["image"] == "D7-017"
    assert saved["n"] == 3
    assert saved["params"] == {"ranges": [dataclasses.asdict(HCTRange(3, 60))]}


def test_record_names_the_versions(tmp_path):
    meta = record(write(tmp_path))["agglpy"]
    assert meta["version"]  # agglpy's own
    assert list(meta["versions"]) == [
        "numpy",
        "pandas",
        "opencv",
        "scipy",
        "scikit-image",
        "matplotlib",
    ]
    assert meta["versions"]["numpy"] == np.__version__
    assert meta["versions"]["pandas"] == pd.__version__
    exported = datetime.fromisoformat(meta["exported"])
    assert exported.tzinfo is not None  # local time with its offset


def test_record_lists_the_tables(tmp_path):
    meta = record(write(tmp_path))["agglpy"]
    assert meta["tables"] == [
        {"name": "particles", "rows": 3, "written": "xlsx"},
        {"name": "summary", "rows": 2, "written": "xlsx"},
    ]


def test_info_sheet_is_the_record_flattened(tmp_path):
    params = HCTParams(ranges=(HCTRange(3, 60),))
    out = write(tmp_path, record={"params": {"hct": params}, "empty": []})
    rows = sheet_rows(out, "info")
    assert rows[0] == ("key", "value")
    assert ("params.hct.ranges.0.d_min", 3) in rows
    assert ("empty", "[]") in rows
    keys = [key for key, _ in rows[1:]]
    assert keys[0] == "agglpy.version"


def test_record_reads_back_unicode(tmp_path):
    out = write(tmp_path, record={"sample": "Fe₂O₃ µm"})
    assert record(out)["sample"] == "Fe₂O₃ µm"
    assert "Fe₂O₃" in (out / "record.yml").read_text(encoding="utf-8")


@pytest.mark.parametrize(
    "value", [Path("a"), {1, 2}, object(), pd.Series([1]), {3: "a"}]
)
def test_record_values_must_be_plain(tmp_path, value):
    with pytest.raises(TypeError, match="record"):
        write(tmp_path, record={"x": value})
    assert not (tmp_path / "csv").exists()


def test_record_key_agglpy_is_reserved(tmp_path):
    with pytest.raises(ParamsError, match="agglpy"):
        write(tmp_path, record={"agglpy": 1})


# ---------------------------------------------------------------------
# An earlier export in the folder
# ---------------------------------------------------------------------


def test_an_earlier_export_is_refused(tmp_path):
    write(tmp_path)
    before = (tmp_path / "record.yml").read_bytes()
    with pytest.raises(OutputExistsError, match="overwrite=True") as info:
        write(tmp_path, tables={"other": SUMMARY})
    assert "results.xlsx" in str(info.value)
    assert (tmp_path / "record.yml").read_bytes() == before
    assert not (tmp_path / "csv" / "other.csv").exists()


def test_overwrite_replaces_only_agglpys_files(tmp_path):
    write(tmp_path)
    (tmp_path / "notes.txt").write_text("mine")
    (tmp_path / "csv" / "stale.csv").write_text("old")
    write(tmp_path, tables={"other": SUMMARY}, overwrite=True)
    assert sorted(p.name for p in (tmp_path / "csv").iterdir()) == [
        "other.csv"
    ]
    assert (tmp_path / "notes.txt").read_text() == "mine"
    assert sheets(tmp_path) == ["other", "info"]


def test_any_one_file_of_an_export_counts(tmp_path):
    (tmp_path / "record.yml").write_text("x: 1\n")
    with pytest.raises(OutputExistsError, match=r"record\.yml"):
        write(tmp_path)


def test_a_locked_file_stops_the_overwrite(tmp_path, monkeypatch):
    write(tmp_path)
    before = {p: p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    rename = Path.rename

    def locked(self, target):
        if self.name == "results.xlsx":  # open in Excel
            raise PermissionError(13, "in use", str(self))
        return rename(self, target)

    monkeypatch.setattr(Path, "rename", locked)
    with pytest.raises(OutputLockedError, match=r"results\.xlsx"):
        write(tmp_path, tables={"other": SUMMARY}, overwrite=True)
    after = {p: p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    assert after == before  # nothing changed


def test_output_errors_are_os_errors(tmp_path):
    write(tmp_path)
    with pytest.raises(FileExistsError):
        write(tmp_path)
